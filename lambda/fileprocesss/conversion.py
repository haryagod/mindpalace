import boto3
import os
import io
import json
import zipfile
from datetime import datetime
from requests_aws4auth import AWS4Auth
from opensearchpy import OpenSearch, RequestsHttpConnection
import re

# ------------------ ENV VARS ------------------
region = os.environ["AWS_REGION"]
opensearch_endpoint = os.environ["OPENSEARCH_ENDPOINT"]
index_name = os.environ["INDEX_NAME"]
metadata_table_name = os.environ["METADATA_TABLE"]
bucket_name = os.environ["BUCKET_NAME"]

session = boto3.Session()
credentials = session.get_credentials().get_frozen_credentials()

awsauth = AWS4Auth(
    credentials.access_key,
    credentials.secret_key,
    region,
    "es",
    session_token=credentials.token,
)

client = OpenSearch(
    hosts=[{"host": opensearch_endpoint, "port": 443}],
    http_auth=awsauth,
    use_ssl=True,
    verify_certs=True,
    timeout=60,
    max_retries=3,
    connection_class=RequestsHttpConnection,
)

s3 = boto3.client("s3")
bedrock = boto3.client("bedrock-runtime")
dynamodb = boto3.resource("dynamodb")
metadata_table = dynamodb.Table(metadata_table_name)

SUPPORTED_FORMATS = [
    "txt", "pdf", "docx", "pptx", "xlsx",
    "json", "html", "xml", "zip"
]

TITAN_V2_DIM = 1024

# ------------------ INPUT CONVERTERS ------------------
def convert_txt(data): return data.decode(errors="ignore")
def convert_pdf(data):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
def convert_docx(data):
    from docx import Document
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
def convert_pptx(data):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    txt = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"): txt.append(shape.text)
    return "\n".join(txt)
def convert_xlsx(data):
    import pandas as pd
    return pd.read_excel(io.BytesIO(data), dtype=str).to_csv(index=False)
def convert_json(data): return json.dumps(json.loads(data), indent=2)
def convert_xml(data): return data.decode(errors="ignore")
def convert_html(data):
    import html2text
    return html2text.html2text(data.decode(errors="ignore"))
def convert_zip(data):
    out = []
    with zipfile.ZipFile(io.BytesIO(data), "r") as z:
        for name in z.namelist():
            if name.endswith("/"): continue
            ext = name.split(".")[-1].lower()
            if ext not in SUPPORTED_FORMATS: continue
            try:
                out.append((name, ext, to_text(ext, z.read(name))))
            except Exception: continue
    return out

INPUT_CONVERTERS = {
    "txt": convert_txt, "pdf": convert_pdf, "docx": convert_docx,
    "pptx": convert_pptx, "xlsx": convert_xlsx, "json": convert_json,
    "xml": convert_xml, "html": convert_html, "zip": convert_zip
}

def to_text(ext, data):
    if ext not in INPUT_CONVERTERS: raise ValueError(f"Unsupported format: {ext}")
    return INPUT_CONVERTERS[ext](data)

# ------------------ OUTPUT GENERATORS ------------------
def output_txt(text): return text.encode(), "text/plain"
def output_json(text): return json.dumps({"text": text}, indent=2).encode(), "application/json"
def output_html(text): return f"<html><body><pre>{text}</pre></body></html>".encode(), "text/html"
def output_xml(text): return f"<root><content><![CDATA[{text}]]></content></root>".encode(), "application/xml"
def output_docx(text):
    from docx import Document
    doc = Document()
    for line in text.split("\n"): doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
def output_pdf(text):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.split("\n"): pdf.multi_cell(0, 7, line)
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue(), "application/pdf"

OUTPUT_GENERATORS = {
    "txt": output_txt, "json": output_json, "html": output_html,
    "xml": output_xml, "docx": output_docx, "pdf": output_pdf
}

# ------------------ EMBEDDING ------------------
def get_embedding(text):
    text = text[:8000]
    resp = bedrock.invoke_model(
        modelId="amazon.titan-embed-text-v2:0",
        accept="application/json",
        contentType="application/json",
        body=json.dumps({"inputText": text}),
    )
    body = json.loads(resp["body"].read())
    emb = body.get("embedding")
    return emb if isinstance(emb, list) else None

# ------------------ OPENSEARCH ------------------
def ensure_index(dimension=TITAN_V2_DIM):
    if client.indices.exists(index=index_name): return False
    body = {
        "settings": {"index": {"knn": True}, "number_of_shards": 1, "number_of_replicas": 1},
        "mappings": {
            "properties": {
                "fileKey": {"type": "keyword"},
                "content": {"type": "text"},
                "embedding": {
                    "type": "knn_vector",
                    "dimension": dimension,
                    "method": {"name": "hnsw", "space_type": "l2", "engine": "nmslib"},
                },
            }
        }
    }
    client.indices.create(index=index_name, body=body)
    return True

def index_document(doc_id, doc):
    if doc.get("embedding") is None: doc.pop("embedding", None)
    return client.index(index=index_name, id=doc_id, body=doc)

def delete_index():
    if client.indices.exists(index=index_name):
        client.indices.delete(index=index_name)
        return f"Index '{index_name}' deleted."
    return f"Index '{index_name}' does not exist."

def search_embedding(query_emb, k=5):
    body = {
        "size": k,
        "query": {
            "knn": {
                "embedding": {
                    "vector": query_emb,
                    "k": k
                }
            }
        }
    }
    res = client.search(index=index_name, body=body)
    hits = [{"id": h["_id"], "score": h["_score"], "source": h["_source"]} for h in res["hits"]["hits"]]
    return hits

def generate_presigned_upload_url(filename: str, expiration_minutes: int = 15):
    key = safe_s3_key(filename)
    url = s3.generate_presigned_url(
        ClientMethod="put_object",
        Params={"Bucket": bucket_name, "Key": key},
        ExpiresIn=expiration_minutes * 60
    )
    return {"key": key, "url": url}

def safe_s3_key(filename: str) -> str:
    """
    Generate a safe S3 key from a filename.
    
    - Converts spaces and unsafe characters to underscores
    - Preserves file extension
    """
    if not filename:
        raise ValueError("Filename cannot be empty")

    # Split base name and extension
    if "." in filename:
        base, ext = ".".join(filename.split(".")[:-1]), filename.split(".")[-1]
        ext = ext.lower()
    else:
        base, ext = filename, ""
    
    # Replace spaces and unsafe characters with underscores
    safe_base = re.sub(r'[^A-Za-z0-9_\-]', '_', base)
    
    # Combine
    safe_key = f"{safe_base}.{ext}" if ext else safe_base
    return safe_key
# ------------------ LAMBDA HANDLER ------------------
def lambda_handler(event, ctx):
    try:
        print("event", str(event) )
        # Detect S3 trigger
        if "Records" in event and "s3" in event["Records"][0]:
            record = event["Records"][0]["s3"]
            bucket = record["bucket"]["name"]
            key = record["object"]["key"]
            mode = "uploadcompleted"
            output_format = None
        elif "body" in event:
            body = event["body"]
            if isinstance(body, str):
                body = json.loads(body)  # convert JSON string to dict
            mode = event.get("path")
            output_format = body.get("outputFormat", "txt")
            search_query = body.get("query")
            key = body.get("key", "")
            bucket = body.get("bucket", bucket_name)
            k = 5

        else:
            bucket = event.get("bucket")
            key = event.get("key")
            mode = event.get("mode", "convert")
            output_format = event.get("outputFormat", "txt")
            search_query = event.get("query")  # For search mode
            print("query:",search_query)
            k = int(event.get("k", 5))

        ext = key.split(".")[-1].lower() if key else None
        if mode == "/upload" and key and ext not in SUPPORTED_FORMATS:
            return {
                "statusCode": 400,
                "headers": {"Content-Type": "application/json"},
                "body": json.dumps({"success": False, "error": f"Unsupported input format: {ext}"})
            }

        responses = []

        # if mode == "search" and search_query:
        #     query_emb = get_embedding(search_query)
        #     hits = search_embedding(query_emb, k)
        #     responses = hits
        #     api_response = {"success": True, "results": responses}
        if mode == "/search" and search_query:
            query_emb = get_embedding(search_query)
            hits = search_embedding(query_emb, k)

            print(str(hits))
    # 3️⃣ Strip embeddings from hits for API response
            hits_cleaned = [
                {
                "id": h.get("id", None),
                "score": h.get("score", None),
                "source": { "fileKey": h.get("source", {}).get("fileKey"), "content": h.get("source", {}).get("content")}
                }
                for h in hits
            ] 

    # 4️⃣ Combine retrieved text for AI context
            context_texts = [h["source"]["content"] for h in hits_cleaned if h["source"].get("content")]
            context_combined = "\n---\n".join(context_texts)

    # 5️⃣ Build AI prompt
            prompt = f"""
            You are a personal assistant for a private knowledge repository called 'MindPalace'.
            Your job is to answer user questions accurately using the context provided, which may include:

            1. Personal information (e.g., names, birthdates, PAN numbers, documents about relatives, etc.).
            2. Research notes and coding snippets (for guidance, R&D, or problem-solving).
            3. General text documents, PDFs, images (transcribed to text), or any structured notes.

            Instructions:
            - Only use the information provided in the context. Do not assume facts that are not present.
            - If the answer is not found in the context, politely say you do not know.
            - Provide concise, clear answers in a conversational style.
            - For coding-related questions, use the context to produce code examples or explanations.
            - Respect privacy and confidentiality; do not fabricate personal data.

            Context (from your documents):
            {context_combined}

            User Question:
            {search_query}

            Answer the question as a short, helpful chat response.
            """

    # 6️⃣ Invoke Bedrock LLM
            response = bedrock.invoke_model(
                modelId="amazon.titan-text-express-v1",
                contentType="application/json",
                accept="application/json",
                body=json.dumps({"inputText": prompt}),
            )
            print("context_combined", context_combined)
            resp_body = response["body"].read()
            resp_json = json.loads(resp_body.decode("utf-8"))
            print("Raw Bedrock body:", resp_body)
            ai_answer = resp_json.get("results", [{}])[0].get("outputText", "")
            api_response = {"success": True, "results": hits_cleaned, "answer": ai_answer}
        elif mode == "/upload":
            if not key:
                return {
                "statusCode": 400,
                "body": json.dumps({"success": False, "error": "Missing 'key' for upload"})
            }
    
            presigned_url = generate_presigned_upload_url(key)
            api_response = {"success": True, "uploadUrl": presigned_url}

        else:
            # handle file conversion or upload
            data = s3.get_object(Bucket=bucket, Key=key)["Body"].read()
            files = convert_zip(data) if ext == "zip" else [(key, ext, to_text(ext, data))]

            for fname, fext, text in files:
                if mode == "/convert":
                    out_bytes, mime = OUTPUT_GENERATORS[output_format](text)
                    out_key = f"converted/{key}/{fname}.{output_format}" if ext=="zip" else f"converted/{fname}.{output_format}"
                    s3.put_object(Bucket=bucket, Key=out_key, Body=out_bytes, ContentType=mime)
                    metadata_table.put_item(Item={
                        "fileKey": fname, "bucket": bucket,
                        "s3InputPath": f"s3://{bucket}/{key}",
                        "outputKey": out_key,
                        "s3OutputPath": f"s3://{bucket}/{out_key}",
                        "inputFormat": fext, "outputFormat": output_format,
                        "mimeType": mime, "updatedAt": datetime.utcnow().isoformat()
                    })
                    responses.append({"file": fname, "mode": "convert", "outputS3Key": out_key})
                elif mode == "uploadcompleted":
                    ensure_index()
                    embedding = get_embedding(text)
                    doc = {"fileKey": fname, "content": text[:5000], "embedding": embedding}
                    index_document(fname.replace("/", "_"), doc)
                    responses.append({"file": fname, "mode": "uploadcompleted", "embeddingStored": embedding is not None})

            api_response = {"success": True, "filesProcessed": len(files), "details": responses}

        # ---------------- Return proper API Gateway format if not S3 event ----------------
        if "Records" in event:
            return api_response  # S3 trigger
        else:
            return {"statusCode": 200, "headers": {"Content-Type": "application/json", "Access-Control-Allow-Origin": "*"}, "body": json.dumps(api_response)}

    except Exception as ex:
        print("ERROR:", str(ex))
        if "Records" in event:
            return {"success": False, "error": str(ex)}
        return {"statusCode": 500, "headers": {"Content-Type": "application/json", "Access-Control-Allow-Origin": "*"}, "body": json.dumps({"success": False, "error": str(ex)})}
